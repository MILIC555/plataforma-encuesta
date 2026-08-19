import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from sqlalchemy.orm import Session
from app.db.models.encuesta import Encuesta
from app.db.models.curso import Curso
from app.db.models.pregunta import Pregunta
from app.db.models.respuesta import Respuesta
from app.analytics.analitica import obtener_kpis_tablero, obtener_desglose_preguntas, obtener_insights_ia, _aplicar_filtros


def generar_excel_encuestas(
    db: Session,
    id_curso: int = None,
    fecha_inicio = None,
    fecha_fin = None,
) -> bytes:
    """Genera un archivo Excel (.xlsx) completo con resumen y detalle de respuestas."""
    wb = Workbook()

    fuente_encabezado = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    relleno_encabezado = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    relleno_sub_encabezado = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
    fuente_titulo = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    fuente_negrita = Font(name="Calibri", size=11, bold=True)
    alinear_centro = Alignment(horizontal="center", vertical="center")
    borde_fino = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    # -------------------------------------------------------------
    # HOJA 1: RESUMEN EJECUTIVO
    # -------------------------------------------------------------
    ws_resumen = wb.active
    ws_resumen.title = "Resumen Ejecutivo"
    ws_resumen.views.sheetView[0].showGridLines = True

    kpis = obtener_kpis_tablero(db, id_curso=id_curso, fecha_inicio=fecha_inicio, fecha_fin=fecha_fin)
    preguntas = obtener_desglose_preguntas(db, id_curso=id_curso, fecha_inicio=fecha_inicio, fecha_fin=fecha_fin)

    curso_obj = db.query(Curso).filter(Curso.id == id_curso).first() if id_curso else None
    nombre_curso = curso_obj.nombre if curso_obj else "Todos los Cursos"

    ws_resumen.cell(row=1, column=1, value="CAMPUS CÓRDOBA — INFORME EJECUTIVO DE ENCUESTAS").font = fuente_titulo
    ws_resumen.cell(row=2, column=1, value=f"Curso: {nombre_curso}").font = fuente_negrita

    # Tabla de KPIs
    ws_resumen.cell(row=4, column=1, value="Indicador Clave").fill = relleno_encabezado
    ws_resumen.cell(row=4, column=1).font = fuente_encabezado
    ws_resumen.cell(row=4, column=2, value="Valor").fill = relleno_encabezado
    ws_resumen.cell(row=4, column=2).font = fuente_encabezado

    filas_kpi = [
        ("Total de Encuestas Procesadas", kpis.get("total_surveys", 0)),
        ("Calificación Promedio General (1 a 10)", round(kpis.get("overall_average", 0.0), 2)),
        ("NPS Score (Recomendación)", f"{kpis.get('nps', 0.0):+0.1f}"),
        ("Índice de Satisfacción General (Q06 >= 8)", f"{kpis.get('satisfaction_pct', 0.0):0.1f}%"),
        ("Porcentaje de Promotores (Q08)", f"{kpis.get('promoters_pct', 0.0):0.1f}%"),
        ("Porcentaje de Detractores (Q08)", f"{kpis.get('detractors_pct', 0.0):0.1f}%"),
    ]

    r_idx = 5
    for etiqueta, val in filas_kpi:
        ws_resumen.cell(row=r_idx, column=1, value=etiqueta).border = borde_fino
        cell_v = ws_resumen.cell(row=r_idx, column=2, value=val)
        cell_v.border = borde_fino
        cell_v.alignment = alinear_centro
        r_idx += 1

    # Tabla de Preguntas Cuantitativas
    r_idx += 2
    ws_resumen.cell(row=r_idx, column=1, value="Cód.").fill = relleno_sub_encabezado
    ws_resumen.cell(row=r_idx, column=1).font = fuente_encabezado
    ws_resumen.cell(row=r_idx, column=2, value="Dimensión / Aspecto").fill = relleno_sub_encabezado
    ws_resumen.cell(row=r_idx, column=2).font = fuente_encabezado
    ws_resumen.cell(row=r_idx, column=3, value="Pregunta Oficial").fill = relleno_sub_encabezado
    ws_resumen.cell(row=r_idx, column=3).font = fuente_encabezado
    ws_resumen.cell(row=r_idx, column=4, value="Promedio (1-10)").fill = relleno_sub_encabezado
    ws_resumen.cell(row=r_idx, column=4).font = fuente_encabezado

    r_idx += 1
    for q in preguntas:
        ws_resumen.cell(row=r_idx, column=1, value=f"Q0{q['question_number']}").border = borde_fino
        ws_resumen.cell(row=r_idx, column=2, value=q["short_label"]).border = borde_fino
        ws_resumen.cell(row=r_idx, column=3, value=q["description"]).border = borde_fino
        cell_avg = ws_resumen.cell(row=r_idx, column=4, value=round(q["average"], 2))
        cell_avg.border = borde_fino
        cell_avg.alignment = alinear_centro
        cell_avg.font = fuente_negrita
        r_idx += 1

    ws_resumen.column_dimensions["A"].width = 30
    ws_resumen.column_dimensions["B"].width = 28
    ws_resumen.column_dimensions["C"].width = 65
    ws_resumen.column_dimensions["D"].width = 18

    # -------------------------------------------------------------
    # HOJA 2: DETALLE DE ENCUESTAS INDIVIDUALES
    # -------------------------------------------------------------
    ws_detalle = wb.create_sheet(title="Detalle Respuestas")
    ws_detalle.views.sheetView[0].showGridLines = True

    encabezados_detalle = [
        "ID Respuesta", "Curso", "Fecha Envío", "Año", "Mes", "Semana",
        "Q01 Aula Virtual", "Q02 Contenidos", "Q03 Material", "Q04 Tutoría",
        "Q05 Administración", "Q06 Satisfacción", "Q07 Expectativas", "Q08 Recomendación",
        "Q09 Comentario Abierto", "Tema IA", "Sentimiento IA"
    ]

    for col_num, h_title in enumerate(encabezados_detalle, 1):
        c = ws_detalle.cell(row=1, column=col_num, value=h_title)
        c.font = fuente_encabezado
        c.fill = relleno_encabezado
        c.alignment = alinear_centro

    # Consultar todas las encuestas y sus respuestas
    query_encuestas = _aplicar_filtros(db.query(Encuesta), id_curso=id_curso, fecha_inicio=fecha_inicio, fecha_fin=fecha_fin)
    encuestas = query_encuestas.order_by(Encuesta.fecha_envio.desc()).all()

    p_objs = {p.id: p.nro_pregunta for p in db.query(Pregunta).all()}

    row_num = 2
    for s in encuestas:
        num_vals = {i: None for i in range(1, 9)}
        comment_text = ""
        ai_topic = ""
        ai_sent = ""

        for r in s.respuestas:
            q_num = p_objs.get(r.id_pregunta)
            if q_num and 1 <= q_num <= 8:
                num_vals[q_num] = r.valor_numerico
            elif q_num == 9:
                comment_text = r.valor_texto or ""
                ai_topic = r.ai_tema or ""
                ai_sent = r.ai_sentimiento or ""

        row_data = [
            s.id_respuesta_origen,
            s.curso.nombre if s.curso else "",
            s.fecha_envio.strftime("%Y-%m-%d %H:%M:%S") if s.fecha_envio else "",
            s.periodo_anio,
            s.periodo_mes,
            s.periodo_semana,
            num_vals[1],
            num_vals[2],
            num_vals[3],
            num_vals[4],
            num_vals[5],
            num_vals[6],
            num_vals[7],
            num_vals[8],
            comment_text,
            ai_topic,
            ai_sent,
        ]

        for col_idx, val in enumerate(row_data, 1):
            cell = ws_detalle.cell(row=row_num, column=col_idx, value=val)
            cell.border = borde_fino
            if col_idx in [1, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 17]:
                cell.alignment = alinear_centro

        row_num += 1

    ws_detalle.column_dimensions["A"].width = 14
    ws_detalle.column_dimensions["B"].width = 35
    ws_detalle.column_dimensions["C"].width = 20
    ws_detalle.column_dimensions["O"].width = 50
    ws_detalle.column_dimensions["P"].width = 20
    ws_detalle.column_dimensions["Q"].width = 16

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


# Alias de compatibilidad
build_survey_excel = generar_excel_encuestas
